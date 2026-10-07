from pathlib import Path
import pytest

ROOT = Path(__file__).resolve().parents[1]


def test_office_builder_output_paths_are_home_confined_and_non_overwriting():
    source = (ROOT / "actions" / "office_builder.py").read_text(encoding="utf-8")
    assert "Office output path must remain inside the user's home directory." in source
    assert "Office output path may not contain symlinked components." in source
    assert "Refusing to overwrite existing Office output without overwrite=True" in source
    assert "overwrite=bool(parameters.get(" in source


def test_office_builder_implicit_outputs_avoid_collisions():
    source = (ROOT / "actions" / "office_builder.py").read_text(encoding="utf-8")
    assert "if overwrite or not base.exists():" in source
    assert "candidate = DEFAULT_OUTPUT_DIR / f" in source



def test_real_presentation_is_structurally_readable_after_save(tmp_path):
    from actions import office_builder as office

    output = tmp_path / "deck.pptx"
    office.create_presentation({
        "title": "Structural Verification",
        "slides": [{"title": "Test", "bullets": ["Body"]}],
        "output_path": str(output),
        "auto_open": False,
    })
    Presentation, *_ = office._import_pptx()
    reopened = Presentation(output)
    assert len(reopened.slides) == 2


def test_real_spreadsheet_is_structurally_readable_after_save(tmp_path):
    from actions import office_builder as office
    from openpyxl import load_workbook

    output = tmp_path / "sheet.xlsx"
    office.create_spreadsheet({
        "title": "Structural Verification",
        "sheets": [{"name": "Sheet1", "headers": ["Item", "Value"], "rows": [["A", 1]]}],
        "output_path": str(output),
        "auto_open": False,
    })
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Sheet1"]
        assert workbook["Sheet1"]["A2"].value == "A"
        assert workbook["Sheet1"]["B2"].value == 1
    finally:
        workbook.close()


def test_presentation_save_failure_cannot_claim_created(monkeypatch, tmp_path):
    from actions import office_builder as office

    Presentation, RGBColor, MSO_SHAPE, Inches, Pt = office._import_pptx()

    class NoSavePresentation(Presentation):
        def save(self, _path):
            return None

    monkeypatch.setattr(
        office,
        "_import_pptx",
        lambda: (NoSavePresentation, RGBColor, MSO_SHAPE, Inches, Pt),
    )
    with pytest.raises(RuntimeError, match="save could not be verified"):
        office.create_presentation({
            "title": "Verification Test",
            "slides": [{"title": "Test", "bullets": ["Body"]}],
            "output_path": str(tmp_path / "deck.pptx"),
            "auto_open": False,
        })


def test_spreadsheet_save_failure_cannot_claim_created(monkeypatch, tmp_path):
    from actions import office_builder as office

    Workbook, BarChart, LineChart, PieChart, Reference, Alignment, Font, PatternFill, get_column_letter = office._import_openpyxl()

    class NoSaveWorkbook(Workbook):
        def save(self, _path):
            return None

    monkeypatch.setattr(
        office,
        "_import_openpyxl",
        lambda: (
            NoSaveWorkbook, BarChart, LineChart, PieChart, Reference,
            Alignment, Font, PatternFill, get_column_letter,
        ),
    )
    with pytest.raises(RuntimeError, match="save could not be verified"):
        office.create_spreadsheet({
            "title": "Verification Test",
            "output_path": str(tmp_path / "sheet.xlsx"),
            "auto_open": False,
        })
